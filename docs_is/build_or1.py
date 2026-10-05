#!/usr/bin/env python
"""
Online Resource 1 (supplementary PDF) and Online Resource 2 (per-section XLSX) for the IS submission.

    python docs_is/build_or1.py
"""
from __future__ import annotations

import csv
import json
import subprocess
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts" / "is_figures"))
from isdata import EXTERNAL_ORDER, PRIMARY_ORDER, REPLICATION_ORDER, VAL, j, ledger, bh  # noqa: E402
from build_is import LATEX_PREAMBLE, fill, keep_tables_together, load_facts  # noqa: E402

DOCS = ROOT / "docs_is"
BUILD = DOCS / "build"
FIGS = ROOT / "results" / "figures" / "is" / "supplement"
FIGM = ROOT / "results" / "figures" / "is"


def md_table(header, rows, align=None, widths=None):
    """Pipe table; `widths` (relative dash counts) sets the column widths pandoc uses when cells wrap."""
    align = align or ["l"] * len(header)
    widths = widths or [2] * len(header)
    sep = ["|:" + "-" * w if a == "l" else "|" + "-" * w + ":" if a == "r" else "|:" + "-" * w + ":"
           for a, w in zip(align, widths)]
    out = ["| " + " | ".join(header) + " |", "".join(sep) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(x) for x in r) + " |")
    return "\n".join(out)


def table_s1():
    led = ledger()
    with open(ROOT / "data" / "admission_audit.csv", encoding="utf-8") as f:
        aud = {r["slide_id"]: r for r in csv.DictReader(f)}
    gc = {r["section"]: r for r in j("graph_coverage_table.json")["rows"]}
    rows = []
    for sid in PRIMARY_ORDER + EXTERNAL_ORDER:
        L, A = led[sid], aud[sid]
        plat = "Visium" if L["platform"] == "visium" else "ST (1st gen.)"
        failed = A["failed"].replace(";", ", ") if A["failed"] else "none"
        acc = ACC_SHORT.get(L["accession"], L["accession"])
        rows.append([sid, L["patient"].replace("_", " "), acc, plat, gc[sid]["n_spots"],
                     f"{float(L['median_umi']):,.0f}", A["min_spots"], A["min_median_umi"], failed])
    return md_table(["Section", "Patient", "Accession", "Platform", "Spots", "Median UMI", "Min spots",
                     "Min median UMI", "Criteria not met"], rows, ["l", "l", "l", "l", "r", "r", "r", "r", "l"],
                    [8, 9, 12, 11, 6, 8, 7, 8, 9])


# 10x dataset names are too long for the table; the caption gives the full names
ACC_SHORT = {"V1_Breast_Cancer_Block_A_Section_1": "10x BC-A1", "V1_Breast_Cancer_Block_A_Section_2": "10x BC-A2",
             "V1_Human_Lymph_Node": "10x LN"}


def num(x, nd=3):
    """Fixed decimals with a typographic minus; tiny non-zero values keep one more significant digit."""
    v = float(x)
    txt = f"{v:.{nd}f}"
    if v != 0 and abs(v) < 0.5 * 10 ** -nd:
        txt = f"{v:.{nd + 1}f}"
    return txt.replace("-", "\u2212")


def qfmt(q):
    return f"{q:.4f}" if 0.045 <= q < 0.055 else f"{q:.3f}"


GENE_SETS = [
    ("ECM (curated core matrisome subset)", "COL1A1, COL1A2, COL3A1, COL5A1, COL6A1, COL6A2, FN1, LAMB1, TNC, THBS2, FBN1, VCAN, BGN, LUM, DCN", "[30]"),
    ("CAF", "COL1A1, COL1A2, COL3A1, DCN, LUM, PDGFRB, FAP, THY1, POSTN", "canonical markers"),
    ("Crosslinking", "LOX, LOXL1, LOXL2, LOXL3, PLOD1, PLOD2, TGM2", "[31, 32]"),
    ("Ligand (absorption proxy)", "CD274, PDCD1LG2", "PD-L1/PD-L2"),
    ("Endothelial", "PECAM1, VWF, CDH5, CLDN5, ENG, EGFL7", "canonical markers"),
    ("T/NK", "CD3D, CD3E, CD2, TRAC, CD8A, GZMB, NKG7, IL7R", "canonical markers"),
    ("Malignant, cSCC", "KRT5, KRT14, KRT6A, KRT17, SFN, S100A2", "[4]"),
    ("Malignant, melanoma", "MLANA, PMEL, TYR, DCT, MITF, SOX10, S100B, TYRP1", "[33]"),
    ("Malignant, breast/lymph node (epithelial)", "EPCAM, KRT8, KRT18, KRT19, MUC1", "epithelial markers"),
    ("Hypoxia (covariate only)", "MSigDB HALLMARK_HYPOXIA, v7.1 (200 genes)", "Liberzon et al. 2015"),
]

PARAMS = [
    ("Graph radius", "150 μm (Visium); 300 μm (1st-gen. ST)", "design; sensitivity in Fig. S4"),
    ("Vessel quantile", "0.80 of endothelial score", "design"),
    ("Source quantile", "0.60 of neighbouring T/NK among vessels", "design"),
    ("Malignant quantile", "0.70 of malignant score", "design"),
    ("Core quantile", "0.50 of distance to non-malignant spots", "design"),
    ("$a$, $b$, $c$ (Eq. 1)", "3, 8, 4", "fixed a priori; not fitted to any outcome"),
    ("$r$ (Eq. 2)", "5.5 nm", "IgG hydrodynamic radius ≈ 5.3 nm [38]"),
    ("$\\xi_0$, $\\beta$ (Eq. 2)", "20 nm, 3", "qualitative scale parameters (Section 2.5)"),
    ("$\\lambda$, $g_0$, $g_{\\mathrm{floor}}$ (Eq. 2)", "3, 1, $10^{-6}$", "fixed a priori; $\\lambda=0$ in ablation"),
    ("$K_{d,\\mathrm{eff}}$, absorption weight (Eq. 3)", "0.5, 1", "uncalibrated; sensitivity analyses"),
    ("Surrogates per test", "500 (primary); 100 (calibration)", "p resolution 1/501"),
    ("Structural-null draws", "200 per section per null", "p resolution 1/201"),
    ("Random seeds", "fixed and stored with every output", "reproducibility"),
]


def table_s4():
    gc = {r["section"]: r for r in j("graph_coverage_table.json")["rows"]}
    rows = []
    for sid in PRIMARY_ORDER + EXTERNAL_ORDER:
        r = gc[sid]
        rows.append([sid, r["n_spots"], r["n_edges"], r["n_components"], f"{100 * r['lcc_share']:.1f}",
                     r["n_vessel"], r["n_source"], r["n_sink"], r["source_outside_lcc"], r["sink_outside_lcc"],
                     f"{100 * r['stranded_share_of_source_sink']:.1f}", r["n_cut_edges"],
                     r["filled_inputs"].replace("ag_target", "ligand").replace(";", ", ") or "–"])
    return md_table(["Section", "Spots", "Edges", "Comp.", "LCC %", "Vessels", "Sources", "Sinks",
                     "Src out", "Sink out", "Stranded %", "Cut edges", "Filled inputs"], rows,
                    ["l"] + ["r"] * 11 + ["l"], [8, 6, 7, 6, 6, 7, 7, 6, 5, 5, 8, 6, 9])


def per_section_rows():
    sn = j("spatial_null_check.json")["per_slide"]
    ext = j("ext_validation.json")["per_slide"]
    ns = j("spatial_null_nscore.json")["per_slide"]
    geo = j("geometry_null.json")["per_slide"]
    abl = j("shared_input_spatial_null.json")["per_slide"]
    adj = j("adjustment_robustness.json")["per_slide"]
    s2 = j("s2_matched_selection.json")["per_slide"]
    s2e = j("s2_matched_ext.json")["per_slide"]
    q_s2 = dict(zip(PRIMARY_ORDER, bh([s2[x]["p_vs_in_cut_matched"] for x in PRIMARY_ORDER])))
    led = ledger()
    rows = []
    for sid in PRIMARY_ORDER + EXTERNAL_ORDER:
        if sid in sn:
            r = sn[sid]
            rho, sd, p, q = r["real_rho_partial"], r["null_std"], r["empirical_p_one_sided"], r["bh_padj"]
            s2r, s2q = s2[sid]["ratio_vs_in_cut_matched"], q_s2[sid]
        else:
            r = ext[sid]["spatial_null"]
            rho, sd, p, q = r["real_rho_partial"], r["null_std"], r["empirical_p_one_sided"], r["bh_padj"]
            s2r, s2q = s2e[sid]["ratio_vs_in_cut_matched"], None
        g, a = geo[sid], abl[sid]
        rows.append(dict(section=sid, patient=led[sid]["patient"], arm="external" if sid in EXTERNAL_ORDER else "primary",
                         rho=rho, null_sd=sd, p=p, q=q, p_nscore=ns[sid]["empirical_p_one_sided"],
                         q_nscore=ns[sid]["bh_padj"], geom_null_mean=g["geometry_null_mean"],
                         constr_null_mean=g["construction_null_mean"], p_constr=g["p_vs_construction_null"],
                         excess=g["excess_over_construction_null"], rho_ablated=a["real_rho_partial"],
                         p_ablated=a["empirical_p_one_sided"], rho_spline=adj[sid]["spline_raw"],
                         rho_unadj=adj[sid]["unadjusted"], s2_ratio=s2r, s2_q=s2q,
                         r_ecm_caf=g["input_correlations"]["ecm_caf"]))
    return rows


def table_s5(rows):
    out = []
    for r in rows:
        out.append([r["section"], num(r["rho"]), num(r["null_sd"]), qfmt(r["q"]), qfmt(r["q_nscore"]),
                    num(r["geom_null_mean"]), num(r["constr_null_mean"]), num(r["p_constr"]),
                    num(r["rho_ablated"]), num(r["s2_ratio"])])
    return md_table(["Section", "ρ", "Null SD", "q (raw)", "q (n-score)", "Geom. null", "Constr. null",
                     "p vs constr.", "ρ ablated", "S2 ratio"], out, ["l"] + ["r"] * 9, [8] + [7] * 9)


def _lab(sid):
    return sid.replace("MEL_THR", "LN P").replace("_rep", " r")


def table_s7():
    """Replication cohort, one row per section."""
    r = j("replication_melanoma.json")["per_slide"]
    rows = []
    for sid in REPLICATION_ORDER:
        v = r[sid]
        rows.append([_lab(sid), v["n_nodes"], v["n_source"], v["n_sink"], num(v["real_rho_partial"]),
                     num(v["null_std"]), qfmt(v["bh_padj"]), qfmt(v["bh_padj_nscore"]),
                     num(v["construction_null_mean"]), num(v["p_vs_construction_null"]), num(v["ablated_rho"]),
                     ", ".join(v["filled_keys"]) if v["filled_keys"] else "none"])
    return md_table(["Section", "Spots", "Src", "Sink", "ρ", "Null SD", "q (raw)", "q (n-score)", "Constr. null",
                     "p vs constr.", "ρ ablated", "Filled inputs"], rows,
                    ["l"] + ["r"] * 10 + ["l"], [9, 5, 4, 4, 6, 6, 6, 6, 6, 6, 6, 10])


def table_s8():
    """CODEX validation: comparators, secondary analyses and sensitivity variants."""
    c = j("codex_validation.json")["summary"]

    def ci(v):
        f = (lambda x: f"{x:.3f}" if abs(x) < 0.01 else f"{x:.2f}")
        return f"{f(v[0])} to {f(v[1])}"
    lab = {"log_b_rel": "SPARTA log $B_{\\mathrm{rel}}$ (primary)", "log_b_cell": "SPARTA log $B_{\\mathrm{cell}}$",
           "field_core": "SPARTA $B_{\\mathrm{cell}}$ field at the core",
           "peritumoural_matrix": "Peritumoural matrix density", "global_matrix": "Global matrix mean",
           "stromal_fraction": "Stromal fraction", "contact_enrichment_z": "Tumour–stroma contact enrichment",
           "ripley_l_matrix": "Ripley's $L$ of matrix-rich cells",
           "vessel_tumour_distance": "Vessel-to-tumour distance (post hoc)",
           "core_depth": "Tumour-core depth (post hoc)"}
    rows = []
    for k, v in c["primary_and_comparators"].items():
        if v.get("degenerate"):
            rows.append([lab[k], "constant", "–", "–", "–"])
            continue
        rows.append([lab[k], num(v["rho_core"], 2), ci(v["ci_core"]),
                     num(v["patient"]["rho"], 2), f"{v['patient']['p_one_sided_perm']:.4f}"])
    pa = c["added_value_partial"]
    rows.append(["Partial log $B_{\\mathrm{rel}}$ given composition", num(pa["rho"], 2), ci(pa["ci"]), "–", "–"])
    ph = c["post_hoc"]
    for key, name in (("b_rel_given_intermixing", "Partial given intermixing (post hoc)"),
                      ("b_rel_given_composition_and_intermixing", "Partial given composition and intermixing (post hoc)"),
                      ("b_rel_given_peritumoural_matrix", "Partial given peritumoural matrix (post hoc)"),
                      ("b_rel_given_vessel_tumour_distance", "Partial given vessel-to-tumour distance (post hoc)"),
                      ("b_rel_given_core_depth", "Partial given tumour-core depth (post hoc)"),
                      ("b_rel_given_distances", "Partial given both distances (post hoc)"),
                      ("b_rel_given_composition_and_distances", "Partial given composition and both distances (post hoc)"),
                      ("b_rel_given_all_simple", "Partial given matrix, both distances and intermixing (post hoc)")):
        if key in ph:
            rows.append([name, num(ph[key]["rho"], 2), ci(ph[key]["ci"]), "–", "–"])
    s1 = c["s1_absolute"]
    rows.append(["S1: log $B_{\\mathrm{cell}}$ vs CD8$^+$ core density", num(s1["rho_core"], 2), ci(s1["ci_core"]),
                 num(s1["patient"]["rho"], 2), f"{s1['patient']['p_one_sided_perm']:.4f}"])
    names = {"edge30": "Edges ≤ 30 μm", "edge100": "Edges ≤ 100 μm", "f_asma_only": "Fibroblast score = αSMA",
             "sinks_all_tumour": "All tumour cells as sinks", "outcome_all_t": "Outcome = all T cells",
             "eq1_half": "Eq. 1 weights halved, b = 4, c = 2 (post hoc)",
             "eq1_x1.5": "Eq. 1 weights × 1.5, b = 12, c = 6 (post hoc)",
             "eq1_matrix_only": "Eq. 1 matrix only, b = 12, c = 0 (post hoc)",
             "eq1_fibroblast_only": "Eq. 1 fibroblast only, b = 0, c = 12 (post hoc)"}
    for k, v in c["sensitivity"].items():
        rows.append([f"Sensitivity: {names.get(k, k)} ({v['n_cores']} cores)", num(v["rho_core"], 2), "–",
                     num(v["patient"]["rho"], 2), f"{v['patient']['p_one_sided_perm']:.4f}"])
    return md_table(["Summary or analysis", "Core ρ", "Core 95% CI", "Patient ρ", "One-sided p"], rows,
                    ["l", "r", "r", "r", "r"], [24, 6, 10, 6, 7])


def table_s9():
    """Intervention maps, one row per section (primary, external, replication)."""
    t = j("intervention_targeting.json")["per_section"]
    ir = j("intervention_ranking.json")["per_section"]

    def joint(sid, op, frac):
        r = ir[sid]
        k = max(1, int(round(frac * r["n"])))
        pts = {q["k"]: q["joint_frac_of_anchor"] for q in r["joint_intervention"][op]["points"]}
        return pts.get(k, float("nan"))
    rows = []
    for sid in list(PRIMARY_ORDER) + list(EXTERNAL_ORDER) + list(REPLICATION_ORDER):
        v = t[sid]
        k = v["k20"]
        name = sid.replace("MEL_THR", "LN P").replace("_rep", " r")
        rows.append([name, v["n"], f"{100 * v['cut_frac']:.0f}",
                     f"{100 * v['k50_cell_frac']:.1f}" if v["k50_cell_frac"] else "–",
                     f"{100 * v['k50_mab_frac']:.0f}" if v["k50_mab_frac"] else "–",
                     f"{100 * joint(sid, 'b_cell', 0.01):.0f}", f"{100 * joint(sid, 'b_mab', 0.01):.0f}",
                     f"{100 * k['cell_sparta']:.0f}", f"{100 * k['cell_random_cut']:.0f}",
                     f"{100 * k['cell_density']:.0f}", f"{100 * k['mab_sparta']:.0f}",
                     f"{100 * k['mab_density']:.0f}", num(v["top5pct_jaccard"], 2)])
    return md_table(["Section", "Spots", "Cut %", "k50 cell %", "k50 mAb %", "Top 1%: cell", "Top 1%: mAb",
                     "20: cell map", "20: rand. cut", "20: cell dens.", "20: mAb map", "20: mAb dens.", "Jaccard"],
                    rows, ["l"] + ["r"] * 12, [9, 5, 4, 5, 5, 5, 5, 5, 5, 5, 5, 5, 5])


def table_s10():
    """Operators against simple spatial summaries, one row per section (run_52)."""
    sb = j("simple_baselines_spearman.json")["per_section"]
    rows = []
    for sid in list(PRIMARY_ORDER) + list(EXTERNAL_ORDER) + list(REPLICATION_ORDER):
        sp_, sc_ = sb[sid]["spot"], sb[sid]["section"]
        name = sid.replace("MEL_THR", "LN P").replace("_rep", " r")
        rows.append([name] + [num(sp_[f"{f}__{b}"], 2) for f in ("b_cell_field", "b_mab_field")
                              for b in ("stromal_density", "dist_tumour_boundary", "niche_z")]
                    + [num(sc_["log_b_rel"], 2), num(sc_["stromal_density"], 3),
                       f"{sc_['dist_tumour_boundary']:.0f}", num(sc_["niche_z"], 1)])
    return md_table(["Section", "Cell: dens.", "Cell: dist.", "Cell: niche", "mAb: dens.", "mAb: dist.", "mAb: niche",
                     "log $B_{\\mathrm{rel}}$", "Peri. density", "Vessel dist.", "Contact z"],
                    rows, ["l"] + ["r"] * 10, [9, 5, 5, 5, 5, 5, 5, 5, 6, 6, 5])


def write_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "per_section"
    cols = list(rows[0].keys())
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    gc = j("graph_coverage_table.json")["rows"]
    ws2 = wb.create_sheet("graph_coverage")
    ws2.append(list(gc[0].keys()))
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in gc:
        ws2.append([r.get(k) for k in gc[0].keys()])
    sb = j("synthetic_benchmark.json")
    ws3 = wb.create_sheet("simulation_summary")
    ws3.append(["summary", "spearman_vs_lost_access", "auc_closed_vs_gap05", "auc_closed_vs_band",
                "auc_closed_vs_scattered", "auc_distant_vs_scattered"])
    for k, v in sb["pooled"].items():
        ws3.append([k, v["spearman_vs_lost_access"], v["auc_closed_vs_gap05"], v["auc_closed_vs_band"],
                    v["auc_closed_vs_scattered"], v.get("auc_closedfar_vs_scattered")])
    nc = j("null_calibration_all.json")
    ws4 = wb.create_sheet("null_calibration")
    ws4.append(["section", "n_nodes", "scenario", "fpr_raw_spectrum", "fpr_normal_score", "fpr_point_level", "n_sim"])
    for sid, r in nc["per_slide"].items():
        for scen in ("gaussian", "marginal", "skewed"):
            if scen in r:
                ws4.append([sid, r["n_nodes"], scen, r[scen]["fpr_spectral_05"], r[scen].get("fpr_nscore_05"),
                            r[scen]["fpr_naive_05"], r[scen]["n_sim"]])
        for tau, v in r.get("smoothness", {}).items():
            ws4.append([sid, r["n_nodes"], f"heat_kernel_{tau}", v["fpr_spectral_05"], v.get("fpr_nscore_05"),
                        v["fpr_naive_05"], v["n_sim"]])
    rep = j("replication_melanoma.json")["per_slide"]
    ws5 = wb.create_sheet("replication_cohort")
    keys5 = ["patient", "n_nodes", "n_source", "n_sink", "real_rho_partial", "null_mean", "null_std",
             "empirical_p_one_sided", "bh_padj", "empirical_p_nscore", "bh_padj_nscore", "geometry_null_mean",
             "construction_null_mean", "p_vs_construction_null", "bh_padj_construction",
             "excess_over_construction_null", "ablated_rho", "ablated_p", "bh_padj_ablated",
             "discordant_low_bcell_high_bmab"]
    ws5.append(["section"] + keys5 + ["filled_keys"])
    for sid in REPLICATION_ORDER:
        ws5.append([sid] + [rep[sid][k] for k in keys5] + [", ".join(rep[sid]["filled_keys"])])
    import csv as _csv
    with open(VAL / "codex_cores.csv", encoding="utf-8") as fh:
        cr = list(_csv.reader(fh))
    ws6 = wb.create_sheet("codex_cores")
    for r_ in cr:
        ws6.append(r_)
    it = j("intervention_targeting.json")["per_section"]
    ws7 = wb.create_sheet("intervention_maps")
    keys7 = ["n", "n_cut_nodes", "cut_frac", "frac_zero_cell", "anchor_cell", "anchor_mab", "k50_cell_frac",
             "k50_mab_frac", "top5pct_jaccard"]
    sub7 = ["cell_sparta", "cell_random_cut", "cell_density", "cell_random", "cell_with_mab_ranking",
            "mab_sparta", "mab_density", "mab_random", "mab_with_cell_ranking"]
    ws7.append(["section"] + keys7 + [f"k20_{k}" for k in sub7] + [f"k5pct_{k}" for k in sub7])
    for sid in list(PRIMARY_ORDER) + list(EXTERNAL_ORDER) + list(REPLICATION_ORDER):
        v = it[sid]
        ws7.append([sid] + [v[k] for k in keys7] + [v["k20"][k] for k in sub7] + [v["k5pct"][k] for k in sub7])
    sb = j("simple_baselines_spearman.json")["per_section"]
    ws8 = wb.create_sheet("simple_summaries")
    sk = sorted(next(iter(sb.values()))["spot"].keys())
    ck = sorted(next(iter(sb.values()))["section"].keys())
    ws8.append(["section", "cohort", "patient"] + [f"spot_{k}" for k in sk] + [f"section_{k}" for k in ck])
    for sid in list(PRIMARY_ORDER) + list(EXTERNAL_ORDER) + list(REPLICATION_ORDER):
        r = sb[sid]
        ws8.append([sid, r["cohort"], r["patient"]] + [r["spot"][k] for k in sk] + [r["section"][k] for k in ck])
    de = j("intervention_domain_enrichment.json")["per_section"]
    ws9 = wb.create_sheet("intervention_compartments")
    ws9.append(["section", "cohort", "patient", "n", "budget", "k_used", "compartment", "share_top", "share_all",
                "enrichment", "p_over_hypergeom", "p_under_hypergeom", "share_cut"])
    for sid in list(PRIMARY_ORDER) + list(EXTERNAL_ORDER) + list(REPLICATION_ORDER):
        r = de[sid]
        for key, v in r["top"].items():
            for d in ("tumour", "stroma", "immune"):
                ws9.append([sid, r["cohort"], r["patient"], r["n"], key, v["k_used"], d, v[d]["frac"],
                            r["base_frac"][d], v[d]["enrichment"], v[d]["p_over"], v[d]["p_under"],
                            r["cut_frac"][d]])
    for w in (ws5, ws6, ws7, ws8, ws9):
        for c in w[1]:
            c.font = Font(bold=True)
    for w in (ws, ws2, ws3, ws4, ws5, ws7, ws8, ws9):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(10, min(28, max(len(str(c.value)) for c in col) + 2))
    out = BUILD / "Online_Resource_2_per_section_results.xlsx"
    wb.save(out)
    return out


def main():
    facts = load_facts()
    text = (DOCS / "online_resource_1.md").read_text(encoding="utf-8")
    rows = per_section_rows()
    text = text.replace("<<TABLE_S1>>", table_s1())
    text = text.replace("<<TABLE_S2>>", md_table(["Signature", "Genes", "Source"], GENE_SETS, widths=[12, 22, 8]))
    text = text.replace("<<TABLE_S3>>", md_table(["Parameter", "Value", "Status"], PARAMS, widths=[12, 14, 14]))
    text = text.replace("<<TABLE_S4>>", table_s4())
    text = text.replace("<<TABLE_S5>>", table_s5(rows))
    text = text.replace("<<TABLE_S7>>", table_s7())
    text = text.replace("<<TABLE_S8>>", table_s8())
    text = text.replace("<<TABLE_S9>>", table_s9())
    text = text.replace("<<TABLE_S10>>", table_s10())
    text = fill(text, facts)
    for tag in ("FigS1a_maps", "FigS1b_maps", "FigS2_calibration", "FigS3_domain_scans", "FigS4_radius",
                "FigS5_reproduction", "FigS6_runtime", "FigS7_intervention"):
        text = text.replace(f"<<{tag}>>", str((FIGS / f"{tag}.pdf").as_posix()))
    src = BUILD / "online_resource_1.md"
    src.write_text(text, encoding="utf-8")
    hdr = BUILD / "preamble_or1.tex"
    hdr.write_text(LATEX_PREAMBLE.replace("\\usepackage[running]{lineno}\\linenumbers", "")
                   .replace("\\renewcommand\\linenumberfont{\\normalfont\\tiny\\sffamily\\color{gray}}", "")
                   .replace("\\usepackage{setspace}\\onehalfspacing", "\\usepackage{setspace}\\setstretch{1.1}")
                   + "\\pretocmd{\\section}{\\needspace{8\\baselineskip}}{}{}\n",
                   encoding="utf-8")
    tex = BUILD / "Online_Resource_1.tex"
    subprocess.run(["pandoc", str(src), "-f", "markdown+tex_math_dollars+pipe_tables+implicit_figures",
                    "-t", "latex", "-s", "-H", str(hdr), "-V", "documentclass=article", "-V", "fontsize=10pt",
                    "-o", str(tex)], check=True)
    t = keep_tables_together(tex.read_text(encoding="utf-8").replace("\\usepackage{lmodern}", ""))
    tex.write_text(t, encoding="utf-8")
    for _ in range(2):
        r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex.name], cwd=BUILD,
                           capture_output=True, text=True, errors="replace")
    if r.returncode != 0:
        print(r.stdout[-2500:])
        raise SystemExit("pdflatex failed (OR1)")
    print("OR1:", BUILD / "Online_Resource_1.pdf")
    print("OR2:", write_xlsx(rows))


if __name__ == "__main__":
    main()
